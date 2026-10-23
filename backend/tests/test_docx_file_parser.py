import os
import tempfile
import pytest
import docx
from app.utils.file_parser import FileParser


def test_docx_in_supported_extensions():
    assert '.docx' in FileParser.SUPPORTED_EXTENSIONS
    assert FileParser.is_supported('market_research.docx')
    assert FileParser.is_supported('/path/to/scenario.DOCX')


def test_extract_text_from_docx():
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, 'test_scenario.docx')
        doc = docx.Document()
        doc.add_paragraph('BRAIN Predictive Intelligence scenario definition.')
        doc.add_paragraph('Key stakeholder: Enterprise CEO and Marketing Director.')
        
        # Add table
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = 'Metric'
        table.rows[0].cells[1].text = 'Target'
        table.rows[1].cells[0].text = 'Conversion'
        table.rows[1].cells[1].text = '18.5%'
        
        doc.save(docx_path)

        extracted = FileParser.extract_text(docx_path)
        assert 'BRAIN Predictive Intelligence scenario definition.' in extracted
        assert 'Key stakeholder: Enterprise CEO and Marketing Director.' in extracted
        assert 'Metric | Target' in extracted
        assert 'Conversion | 18.5%' in extracted


def test_extract_from_multiple_with_docx():
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, 'doc1.docx')
        doc = docx.Document()
        doc.add_paragraph('DOCX paragraph content.')
        doc.save(docx_path)

        txt_path = os.path.join(tmpdir, 'doc2.txt')
        with open(txt_path, 'w', encoding='utf-8') as f:
            f.write('Plain text file content.')

        merged = FileParser.extract_from_multiple([docx_path, txt_path])
        assert 'DOCX paragraph content.' in merged
        assert 'Plain text file content.' in merged
        assert '=== 文档 1: doc1.docx ===' in merged
        assert '=== 文档 2: doc2.txt ===' in merged
